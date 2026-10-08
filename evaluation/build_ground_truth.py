#!/usr/bin/env python3
"""
GreenTrust Module 1 Evaluation — Ground Truth Builder

Independently extracts and detects information from test documents
WITHOUT calling the GreenTrust processor. Produces ground_truth.json
for later comparison with processor output.

Usage:
    python build_ground_truth.py                # interactive review
    python build_ground_truth.py --auto-accept  # skip review
"""

import argparse
import io
import json
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Set

import fitz  # PyMuPDF
import pytesseract
from docx import Document as DocxDocument
from PIL import Image


# =========================================================
# CONFIGURATION
# =========================================================

BASE_DIR = Path(__file__).resolve().parent
TEST_DOCS_DIR = BASE_DIR / "test_documents"
OUTPUT_FILE = BASE_DIR / "ground_truth.json"

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".png", ".jpg", ".jpeg"}


# =========================================================
# REGEX PATTERNS (independent of processor)
# =========================================================

DATE_PATTERN = re.compile(
    r"\b("
    r"\d{1,2}[./-]\d{1,2}[./-]\d{2,4}"
    r"|"
    r"\d{4}[./-]\d{1,2}[./-]\d{1,2}"
    r"|"
    r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
    r"[a-z]*\s+\d{1,2}(?:st|nd|rd|th)?,?\s+\d{2,4}"
    r"|"
    r"\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
    r"[a-z]*\s+\d{2,4}"
    r"|"
    r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
    r"[a-z]*\s+\d{4}"
    r"|"
    r"\d{1,2}[/-]\d{4}"
    r")\b",
    re.IGNORECASE,
)

MEASUREMENT_PATTERN = re.compile(
    r"(?<![\w.])"
    r"([+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)"
    r"\s*"
    r"(KL|kL|kl|m³|m3|MWh|kWh|Wh|kW|MW|W|"
    r"liters?|litres?|L|milliliters?|mL|"
    r"cubic\s+(?:meters?|metres?)|"
    r"tonnes?|tons?|kg|g|mg|%|percent|°C|°F|C|F|V|A)"
    r"\b",
    re.IGNORECASE,
)

MONEY_PATTERN = re.compile(
    r"(?<!\w)"
    r"(?P<currency>₹|₽|€|£|\$|USD|EUR|GBP|INR|AED|AUD|CAD|Rs\.?)"
    r"\s*"
    r"(?P<amount>\d{1,3}(?:,\d{3})*(?:\.\d{1,2})?|\d+(?:\.\d{1,2})?)"
    r"|"
    r"(?<!\w)"
    r"(?P<amount_after>\d{1,3}(?:,\d{3})*(?:\.\d{1,2})?|\d+(?:\.\d{1,2})?)"
    r"\s*"
    r"(?P<currency_after>USD|EUR|GBP|INR|AED|AUD|CAD|rupees?|dollars?|euros?|pounds?)"
    r"\b",
    re.IGNORECASE,
)

KEY_VALUE_PATTERN = re.compile(
    r"^\s*"
    r"([A-Za-z][A-Za-z0-9 /()#&._-]{1,60}?)"
    r"\s*(?::|=|\s[-–—]\s|\t+)"
    r"\s*(.+?)\s*$"
)

NOISE_LINE_PATTERN = re.compile(r"^(?:[\W_]|[-_=*]){3,}$")

# Hotel/property name patterns
HOTEL_NAME_PATTERN = re.compile(
    r"(?:hotel|resort|villa|lodge|guest\s*house|homestay|boutique|eco\s*lodge|serviced\s*apartment)"
    r"[\s:]+([A-Z][A-Za-z0-9\s&'.,-]{2,60})",
    re.IGNORECASE,
)

PROPERTY_NAME_KEYS = {
    "hotel_name", "property_name", "customer_name",
    "name", "establishment", "property",
}


# =========================================================
# SUSTAINABILITY TERMS
# =========================================================

ENERGY_TERMS: Set[str] = {
    "energy", "electricity", "power", "solar", "kwh",
    "consumption", "generator", "grid", "renewable", "meter",
}

WATER_TERMS: Set[str] = {
    "water", "rainwater", "harvesting", "wastewater",
    "sewage", "irrigation", "pool", "liters", "meter",
}

WASTE_TERMS: Set[str] = {
    "waste", "recycling", "compost", "biodegradable",
    "landfill", "collection", "disposal", "segregation",
}

LEGAL_TERMS: Set[str] = {
    "license", "policy", "sustainability", "certification",
    "commitment", "environmental", "compliance", "audit",
}

ALL_TERMS = ENERGY_TERMS | WATER_TERMS | WASTE_TERMS | LEGAL_TERMS


# =========================================================
# TEXT CLEANING
# =========================================================

def clean_text(raw: str) -> str:
    if not raw:
        return ""
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    lines = []
    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        if NOISE_LINE_PATTERN.fullmatch(line):
            continue
        lines.append(line)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


# =========================================================
# INDEPENDENT DOCUMENT EXTRACTION
# =========================================================

def extract_pdf(file_path: Path) -> Dict:
    """Extract text from PDF using PyMuPDF with OCR fallback."""
    buf = file_path.read_bytes()
    doc = fitz.open(stream=buf, filetype="pdf")
    pages_text = []

    for page in doc:
        text = page.get_text("text")
        if text.strip():
            pages_text.append(text)
            continue

        # OCR fallback for image-only pages
        pix = page.get_pixmap(dpi=200)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        ocr_text = pytesseract.image_to_string(img)
        pages_text.append(ocr_text)

    page_count = doc.page_count
    doc.close()
    return {"text": "\n".join(pages_text), "page_count": page_count}


def extract_docx(file_path: Path) -> Dict:
    """Extract text from DOCX."""
    doc = DocxDocument(str(file_path))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return {
        "text": "\n".join(parts),
        "page_count": max(1, len(parts) // 45 + 1),
    }


def extract_image(file_path: Path) -> Dict:
    """Extract text from image using Tesseract OCR."""
    img = Image.open(str(file_path))
    text = pytesseract.image_to_string(img)
    return {"text": text, "page_count": 1}


def extract_document(file_path: Path) -> Optional[Dict]:
    """Route to correct extractor based on file extension."""
    ext = file_path.suffix.lower()
    try:
        if ext == ".pdf":
            result = extract_pdf(file_path)
        elif ext == ".docx":
            result = extract_docx(file_path)
        elif ext in (".png", ".jpg", ".jpeg"):
            result = extract_image(file_path)
        else:
            return None
        result["cleaned_text"] = clean_text(result["text"])
        result["file_type"] = ext.lstrip(".")
        return result
    except Exception as e:
        print(f"  [ERROR] Extraction failed for {file_path.name}: {e}")
        return None


# =========================================================
# FIELD DETECTION
# =========================================================

def detect_dates(text: str) -> List[str]:
    return list(dict.fromkeys(
        m.group(0).strip() for m in DATE_PATTERN.finditer(text)
    ))[:25]


def detect_measurements(text: str) -> List[Dict]:
    results = []
    seen = set()
    for m in MEASUREMENT_PATTERN.finditer(text):
        value_text = m.group(1).replace(",", "")
        try:
            value: Any = float(value_text)
            if value.is_integer():
                value = int(value)
        except ValueError:
            continue

        unit = m.group(2)
        key = (value, unit.lower())

        if key in seen:
            continue
        seen.add(key)

        results.append({"value": value, "unit": unit})

    return results[:30]


def detect_financials(text: str) -> List[Dict]:
    results = []
    seen = set()

    currency_map = {
        "₹": "INR", "rs": "INR", "rs.": "INR", "inr": "INR",
        "$": "USD", "usd": "USD",
        "€": "EUR", "eur": "EUR",
        "£": "GBP", "gbp": "GBP",
        "₽": "RUB",
        "aed": "AED", "aud": "AUD", "cad": "CAD",
    }

    for m in MONEY_PATTERN.finditer(text):
        if m.group("currency"):
            raw_currency = m.group("currency")
            amount_str = m.group("amount")
        else:
            raw_currency = m.group("currency_after")
            amount_str = m.group("amount_after")

        try:
            amount = float(amount_str.replace(",", ""))
            if amount.is_integer():
                amount = int(amount)
        except (ValueError, AttributeError):
            continue

        currency = currency_map.get(raw_currency.strip().lower(), raw_currency.upper() if raw_currency.isalpha() else "INR")

        key = (amount, currency)
        if key in seen:
            continue
        seen.add(key)

        results.append({"amount": amount, "currency": currency})

    return results[:30]


def normalize_key(label: str) -> str:
    return re.sub(r"[^a-zA-Z0-9]+", "_", label.strip().lower()).strip("_")


def detect_key_values(text: str) -> Dict[str, str]:
    fields: Dict[str, str] = {}
    ignored_keys = {"amount", "note", "consumption", "demand"}

    for line in text.split("\n"):
        m = KEY_VALUE_PATTERN.match(line)
        if not m:
            continue
        key = normalize_key(m.group(1))
        value = m.group(2).strip(" |-")
        if not key or not value or len(key) < 2 or len(value) > 240:
            continue
        if key in ignored_keys:
            continue
        if key not in fields:
            fields[key] = value
    return fields


def detect_hotel_name(text: str, key_values: Dict[str, str]) -> Optional[str]:
    # Check key-value fields first
    for k in PROPERTY_NAME_KEYS:
        if k in key_values:
            return key_values[k]

    # Search in text
    m = HOTEL_NAME_PATTERN.search(text)
    if m:
        return m.group(1).strip()

    # Look for ALL-CAPS lines that look like hotel names
    for line in text.split("\n"):
        line = line.strip()
        if len(line) > 5 and line.isupper() and any(
            w in line.lower() for w in ["hotel", "resort", "villa", "lodge"]
        ):
            return line

    return None


def find_relevant_terms(text: str) -> List[str]:
    lowered = text.lower()
    found = []
    for term in sorted(ALL_TERMS):
        if re.search(rf"(?<!\w){re.escape(term)}(?!\w)", lowered):
            found.append(term)
    return found


# =========================================================
# BUILD GROUND TRUTH ENTRY
# =========================================================

def build_entry(file_path: Path, extraction: Dict) -> Dict:
    text = extraction["cleaned_text"]

    dates = detect_dates(text)
    measurements = detect_measurements(text)
    financials = detect_financials(text)
    key_values = detect_key_values(text)
    hotel_name = detect_hotel_name(text, key_values)
    relevant_terms = find_relevant_terms(text)

    entry: Dict = {
        "file_name": file_path.name,
        "file_type": extraction["file_type"],
        "dates": dates,
        "measurements": measurements,
        "financials": financials,
        "key_data": {},
        "relevant_terms": relevant_terms,
    }

    if hotel_name:
        entry["key_data"]["hotel_name"] = hotel_name

    # Include other key-value fields (excluding hotel_name to avoid dup)
    for k, v in key_values.items():
        if k not in PROPERTY_NAME_KEYS:
            entry["key_data"][k] = v

    # Include optional expected_text if text was extracted
    if text.strip():
        entry["expected_text"] = text[:5000]  # cap for practical review

    return entry


# =========================================================
# INTERACTIVE REVIEW
# =========================================================

def review_entry(file_name: str, entry: Dict) -> Dict:
    """Show detected values and let user edit or accept."""
    print("\n" + "=" * 60)
    print(f"  {file_name}")
    print("=" * 60)

    print(f"\n  Dates:          {entry['dates']}")
    print(f"  Measurements:   {entry['measurements']}")
    print(f"  Financials:     {entry['financials']}")
    print(f"  Key data:       {entry['key_data']}")
    print(f"  Relevant terms: {entry['relevant_terms']}")

    while True:
        choice = input("\n  [A]ccept / [E]dit / [S]kip? (a/e/s): ").strip().lower()
        if choice in ("a", ""):
            return entry
        elif choice == "s":
            return None
        elif choice == "e":
            entry = edit_interactive(entry)
            return entry


def edit_interactive(entry: Dict) -> Dict:
    """Simple JSON-based editing."""
    print("\n  Current entry (JSON):")
    print(json.dumps(entry, indent=2, ensure_ascii=False))
    print("\n  Paste replacement JSON, or press Enter to keep current:")
    lines = []
    try:
        while True:
            line = input()
            if not line:
                break
            lines.append(line)
    except EOFError:
        pass

    if lines:
        try:
            new_entry = json.loads("\n".join(lines))
            print("  Updated.")
            return new_entry
        except json.JSONDecodeError as e:
            print(f"  Invalid JSON: {e}. Keeping original.")
    return entry


# =========================================================
# MAIN
# =========================================================

def discover_documents() -> List[Path]:
    """Find all supported test documents."""
    docs = []
    if not TEST_DOCS_DIR.exists():
        return docs
    for f in TEST_DOCS_DIR.rglob("*"):
        if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS:
            docs.append(f)
    return sorted(docs)


def main():
    parser = argparse.ArgumentParser(
        description="Build ground truth for GreenTrust Module 1 evaluation"
    )
    parser.add_argument(
        "--auto-accept",
        action="store_true",
        help="Skip interactive review and auto-accept all detected values",
    )
    args = parser.parse_args()

    print("\n" + "=" * 60)
    print("  GreenTrust Module 1 — Ground Truth Builder")
    print("=" * 60)

    docs = discover_documents()
    if not docs:
        print(f"\n  No documents found in {TEST_DOCS_DIR}")
        print("  Place PDF, DOCX, PNG, or JPG files in:")
        print("    evaluation/test_documents/pdf/")
        print("    evaluation/test_documents/docx/")
        print("    evaluation/test_documents/images/")
        sys.exit(1)

    print(f"\n  Found {len(docs)} document(s):")
    for d in docs:
        print(f"    - {d.relative_to(BASE_DIR)}")

    ground_truth: Dict[str, Dict] = {}
    skipped = 0

    for doc_path in docs:
        rel_name = doc_path.name
        print(f"\n  Processing: {rel_name}")

        extraction = extract_document(doc_path)
        if extraction is None:
            print(f"  [SKIP] Could not extract: {rel_name}")
            skipped += 1
            continue

        entry = build_entry(doc_path, extraction)

        if args.auto_accept:
            ground_truth[rel_name] = entry
            print(f"  [AUTO-ACCEPTED] {len(entry['dates'])} dates, "
                  f"{len(entry['measurements'])} measurements, "
                  f"{len(entry['financials'])} financials, "
                  f"{len(entry['key_data'])} key fields, "
                  f"{len(entry['relevant_terms'])} terms")
        else:
            reviewed = review_entry(rel_name, entry)
            if reviewed is not None:
                ground_truth[rel_name] = reviewed
                print(f"  [ACCEPTED]")
            else:
                print(f"  [SKIPPED]")
                skipped += 1

    # Write ground truth
    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(ground_truth, f, indent=2, ensure_ascii=False)

    print("\n" + "=" * 60)
    print(f"  Ground truth written to: {OUTPUT_FILE}")
    print(f"  Documents: {len(ground_truth)} accepted, {skipped} skipped")
    print("=" * 60 + "\n")


if __name__ == "__main__":
    main()
