"""
GreenTrust Phase 1 — Document Processor (FastAPI)

Extracts text, metadata, and structured JSON from PDF/DOCX/PNG/JPG.

Returns:
1. Simple extraction fields for the frontend.
2. Rich structured JSON for later verification / AI processing.
"""

import hashlib
import io
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Set

import fitz  # PyMuPDF
import pytesseract
from docx import Document as DocxDocument
from fastapi import FastAPI, File, Form, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from PIL import Image


# =========================================================
# APP
# =========================================================

app = FastAPI(
    title="GreenTrust Processor",
    version="1.0.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


# =========================================================
# REGEX PATTERNS
# =========================================================

DATE_PATTERN = re.compile(
    r"\b("
    r"\d{1,2}[./-]\d{1,2}[./-]\d{2,4}"
    r"|"
    r"\d{4}[./-]\d{1,2}[./-]\d{1,2}"
    r"|"
    r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
    r"[a-z]*\s+\d{1,2}(?:st|nd|rd|th)?,?\s+\d{2,4}"
    r"|"
    r"\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
    r"[a-z]*\s+\d{2,4}"
    r"|"
    r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
    r"[a-z]*\s+\d{4}"
    r"|"
    r"\d{1,2}[/-]\d{4}"
    r")\b",
    re.IGNORECASE,
)

MEASUREMENT_PATTERN = re.compile(
    r"(?<![\w.])"
    r"([+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)"
    r"\s*"
    r"(KL|kL|kl|m³|m3|MWh|kWh|Wh|kW|MW|W|"
    r"liters?|litres?|L|milliliters?|mL|"
    r"cubic\s+(?:meters?|metres?)|"
    r"tonnes?|tons?|kg|g|mg|%|percent|°C|°F|C|F|V|A)"
    r"\b",
    re.IGNORECASE,
)

MONEY_PATTERN = re.compile(
    r"(?<!\w)"
    r"(?P<currency>₹|₽|€|£|\$|USD|EUR|GBP|INR|AED|AUD|CAD)"
    r"\s*"
    r"(?P<amount>\d{1,3}(?:,\d{3})*(?:\.\d{1,2})?|\d+(?:\.\d{1,2})?)"
    r"|"
    r"(?<!\w)"
    r"(?P<amount_after>\d{1,3}(?:,\d{3})*(?:\.\d{1,2})?|\d+(?:\.\d{1,2})?)"
    r"\s*"
    r"(?P<currency_after>USD|EUR|GBP|INR|AED|AUD|CAD|rupees?|dollars?|euros?|pounds?)"
    r"\b",
    re.IGNORECASE,
)

KEY_VALUE_PATTERN = re.compile(
    r"^\s*"
    r"([A-Za-z][A-Za-z0-9 /()#&._-]{1,60}?)"
    r"\s*(?::|=|\s[-–—]\s|\t+)"
    r"\s*(.+?)\s*$"
)

NOISE_LINE_PATTERN = re.compile(
    r"^(?:[\W_]|[-_=*]){3,}$"
)


# =========================================================
# RELEVANT TERMS
# =========================================================

ENERGY_TERMS = {
    "energy",
    "electricity",
    "power",
    "solar",
    "kwh",
    "consumption",
    "generator",
    "grid",
    "renewable",
    "meter",
}

WATER_TERMS = {
    "water",
    "rainwater",
    "harvesting",
    "wastewater",
    "sewage",
    "irrigation",
    "pool",
    "consumption",
    "m³",
    "liters",
    "meter",
}


# =========================================================
# TEXT CLEANING
# =========================================================

def clean_text(raw: str) -> str:
    """Normalize extracted text while preserving useful content."""

    if not raw:
        return ""

    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)

    lines = []

    for raw_line in text.split("\n"):
        line = raw_line.strip()

        if not line:
            continue

        if NOISE_LINE_PATTERN.fullmatch(line):
            continue

        # Normalize table separators.
        line = re.sub(r"\s*([|¦])\s*", " | ", line)
        lines.append(line)

    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


# =========================================================
# DATE EXTRACTION
# =========================================================

def detect_dates(text: str) -> List[str]:
    """Return unique dates found in the text."""

    return list(
        dict.fromkeys(
            match.group(0).strip()
            for match in DATE_PATTERN.finditer(text)
        )
    )[:25]


# =========================================================
# MEASUREMENT EXTRACTION
# =========================================================

def detect_measurements(text: str) -> List[str]:
    """Return simple value + unit measurements."""

    return list(
        dict.fromkeys(
            f"{match.group(1)} {match.group(2)}"
            for match in MEASUREMENT_PATTERN.finditer(text)
        )
    )[:30]


def detect_measurement_objects(text: str) -> List[Dict]:
    """Extract measurements as structured objects."""

    results = []

    for match in MEASUREMENT_PATTERN.finditer(text):
        value_text = match.group(1)
        unit = match.group(2)

        try:
            value = float(value_text.replace(",", ""))

            if value.is_integer():
                value = int(value)
        except ValueError:
            value = value_text

        # Look backwards for a possible label.
        preceding = text[max(0, match.start() - 100):match.start()]

        label_match = re.search(
            r"([A-Za-z][A-Za-z0-9 /()#&._-]{2,60})\s*:\s*$",
            preceding,
        )

        label = (
            label_match.group(1).strip()
            if label_match
            else None
        )

        item = {
            "value": value,
            "unit": unit,
        }

        if label:
            item["label"] = label

        results.append(item)

    return deduplicate_measurements(results)[:30]


def deduplicate_measurements(
    measurements: List[Dict],
) -> List[Dict]:
    """Remove duplicate value/unit pairs while preserving useful labels."""

    priority_labels = {
        "sanctioned load",
        "previous reading",
        "current reading",
        "energy consumed",
        "peak demand",
        "power factor",
        "average daily usage",
        "previous month",
        "current month",
    }

    result = {}

    for item in measurements:
        value = item.get("value")
        unit = item.get("unit")
        label = item.get("label", "").strip()
        key = (value, unit)

        if key not in result:
            result[key] = item
            continue

        existing = result[key]
        existing_label = (
            existing.get("label", "").strip().lower()
        )
        current_label = label.lower()

        if (
            current_label in priority_labels
            and existing_label not in priority_labels
        ):
            result[key] = item

    return list(result.values())


# =========================================================
# FINANCIAL EXTRACTION
# =========================================================

def detect_money(text: str) -> List[str]:
    """Return monetary amounts in their original text form."""

    amounts = []

    for match in MONEY_PATTERN.finditer(text):
        if match.group("currency"):
            amounts.append(
                f"{match.group('currency')} {match.group('amount')}"
            )
        else:
            amounts.append(
                f"{match.group('amount_after')} "
                f"{match.group('currency_after')}"
            )

    return list(dict.fromkeys(amounts))[:30]


def extract_financials(text: str) -> List[Dict]:
    """
    Extract monetary values from labelled lines.

    Example:
        1. Energy Consumption Charges
        Amount : Rs. 50,000

    becomes:
        {
            "amount": 50000,
            "currency": "INR",
            "label": "Energy Consumption Charges"
        }
    """

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    financials = []

    money_pattern = re.compile(
        r"^(.*?)\s*:\s*(?:Rs\.?|₹|INR)\s*"
        r"([\d,]+(?:\.\d+)?)",
        re.IGNORECASE,
    )

    ignored_labels = {"amount"}
    current_section = None

    for line in lines:
        # Detect numbered section.
        numbered = re.match(r"^\d+\.\s*(.+)$", line)

        if numbered:
            current_section = numbered.group(1).strip()
            continue

        match = money_pattern.match(line)

        if not match:
            continue

        label = match.group(1).strip()
        amount_text = match.group(2).replace(",", "")

        try:
            amount = float(amount_text)
        except ValueError:
            continue

        if label.lower() in ignored_labels and current_section:
            final_label = current_section
        else:
            final_label = label

        financials.append(
            {
                "amount": amount,
                "currency": "INR",
                "label": final_label,
            }
        )

    return financials


# =========================================================
# KEY-VALUE EXTRACTION
# =========================================================

def normalize_key(label: str) -> str:
    """Convert a human-readable label into a normalized key."""

    key = re.sub(
        r"[^a-zA-Z0-9]+",
        "_",
        label.strip().lower(),
    )

    return key.strip("_")


def detect_key_values(text: str) -> Dict[str, str]:
    """Extract simple key-value fields from labelled lines."""

    fields: Dict[str, str] = {}

    ignored_keys = {
        "amount",
        "consumption",
        "demand",
        "previous_month",
        "current_month",
        "note",
    }

    for line in text.split("\n"):
        match = KEY_VALUE_PATTERN.match(line)

        if not match:
            continue

        key = normalize_key(match.group(1))
        value = match.group(2).strip(" |-")

        if not key or not value:
            continue

        if len(key) < 2 or len(value) > 240:
            continue

        if key in ignored_keys:
            continue

        # Do not overwrite an existing value.
        if key not in fields:
            fields[key] = value

    return fields


# =========================================================
# RELEVANT TERM DETECTION
# =========================================================

def find_relevant_terms(
    text: str,
    terms: Set[str],
) -> List[str]:
    """Find configured sustainability terms in the text."""

    lowered = text.lower()

    found = [
        term
        for term in terms
        if re.search(
            rf"(?<!\w){re.escape(term.lower())}(?!\w)",
            lowered,
        )
    ]

    return sorted(set(found))


# =========================================================
# DOCUMENT EXTRACTION
# =========================================================

def extract_pdf(buf: bytes) -> Dict:
    """Extract PDF text with Tesseract fallback for image-only pages."""

    doc = fitz.open(stream=buf, filetype="pdf")
    pages_text = []
    ocr_used = False

    for page in doc:
        text = page.get_text("text")

        if text.strip():
            pages_text.append(text)
            continue

        pix = page.get_pixmap(dpi=200)
        img = Image.open(io.BytesIO(pix.tobytes("png")))

        ocr_text = pytesseract.image_to_string(img)
        pages_text.append(ocr_text)
        ocr_used = True

    page_count = doc.page_count
    doc.close()

    return {
        "text": "\n".join(pages_text),
        "page_count": page_count,
        "method": (
            "PDF OCR fallback (Tesseract)"
            if ocr_used
            else "PDF text extraction (PyMuPDF)"
        ),
    }


def extract_docx(buf: bytes) -> Dict:
    """Extract DOCX paragraphs and table contents."""

    doc = DocxDocument(io.BytesIO(buf))
    parts = []

    # Paragraphs.
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    # Tables.
    for table in doc.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                parts.append(" | ".join(cells))

    return {
        "text": "\n".join(parts),
        "page_count": max(1, len(parts) // 45 + 1),
        "method": "DOCX paragraph & table extraction",
    }


def extract_image(buf: bytes) -> Dict:
    """Extract text from PNG/JPG/JPEG using Tesseract OCR."""

    img = Image.open(io.BytesIO(buf))
    text = pytesseract.image_to_string(img)

    return {
        "text": text,
        "page_count": 1,
        "method": "Tesseract OCR",
    }


# =========================================================
# STRUCTURED SUMMARY
# =========================================================

def build_summary(
    key_values: Dict[str, str],
    measurements: List[Dict],
    financials: List[Dict],
) -> str:
    """Build a short human-readable summary from extracted fields."""

    customer = (
        key_values.get("customer_name")
        or key_values.get("hotel_name")
        or key_values.get("property_name")
    )

    billing_period = key_values.get("billing_period")

    energy = (
        key_values.get("energy_consumed")
        or key_values.get("total_consumption")
    )

    total_due = None

    for item in financials:
        label = (item.get("label") or "").lower()

        if "total amount due" in label:
            currency = item.get("currency", "INR")
            total_due = f"{currency} {item['amount']:,.2f}"
            break

    parts = []

    if customer:
        parts.append(f"Bill for {customer}")

    if billing_period:
        parts.append(f"covering {billing_period}")

    if energy:
        parts.append(f"with energy consumption of {energy}")

    if total_due:
        parts.append(f"and total amount due of {total_due}")

    if parts:
        return " ".join(parts) + "."

    return "Document containing extracted billing information."


# =========================================================
# STRUCTURED JSON
# =========================================================

def build_structured_data(
    file_name: str,
    file_type: str,
    text: str,
    metadata: Dict,
) -> Dict:
    """Build the rich structured representation returned by the processor."""

    key_values = detect_key_values(text)
    measurements = detect_measurement_objects(text)
    financials = extract_financials(text)
    dates = detect_dates(text)

    meaningful_lines = [
        line.strip()
        for line in text.split("\n")
        if line.strip()
        and not NOISE_LINE_PATTERN.fullmatch(line.strip())
    ]

    key_excerpts = [
        line
        for line in meaningful_lines
        if len(line) > 20
    ][:8]

    relevant = find_relevant_terms(
        text,
        ENERGY_TERMS | WATER_TERMS,
    )

    redundant_keys = {
        "consumption",
        "demand",
        "amount",
        "previous_month",
        "current_month",
    }

    key_values = {
        key: value
        for key, value in key_values.items()
        if key not in redundant_keys
    }

    summary = build_summary(
        key_values,
        measurements,
        financials,
    )

    return {
        "document": {
            "file_name": file_name,
            "file_type": file_type,
            "page_count": metadata["page_count"],
        },
        "summary": summary,
        "key_data": key_values,
        "measurements": measurements,
        "financials": financials,
        "dates": dates,
        "relevant_terms": relevant,
        "key_excerpts": key_excerpts,
        "metadata": {
            "word_count": metadata["word_count"],
            "line_count": metadata["line_count"],
            "extraction_method": metadata["extraction_method"],
            "processed_timestamp": metadata["processed_timestamp"],
        },
    }


# =========================================================
# DOCUMENT PROCESSING ENDPOINT
# =========================================================

@app.post("/process/document")
async def process_document(
    file: UploadFile = File(...),
    document_id: str = Form(None),
    file_type: str = Form(None),
):
    """Extract, clean, structure, and hash an uploaded document."""

    # -----------------------------------------------------
    # 1. Read uploaded file
    # -----------------------------------------------------

    buf = await file.read()

    # -----------------------------------------------------
    # 2. Determine file type
    # -----------------------------------------------------

    ftype = (
        file_type
        or Path(file.filename or "").suffix.lstrip(".")
    ).lower()

    # -----------------------------------------------------
    # 3. Extract document content
    # -----------------------------------------------------

    try:
        if ftype == "pdf":
            extraction = extract_pdf(buf)

        elif ftype == "docx":
            extraction = extract_docx(buf)

        elif ftype in {"png", "jpg", "jpeg"}:
            extraction = extract_image(buf)

            if ftype == "jpeg":
                ftype = "jpg"

        else:
            return JSONResponse(
                status_code=400,
                content={
                    "error": f"Unsupported file type: {ftype}",
                },
            )

    except Exception as exc:
        return JSONResponse(
            status_code=422,
            content={"error": str(exc)},
        )

    # -----------------------------------------------------
    # 4. Clean extracted text
    # -----------------------------------------------------

    cleaned = clean_text(extraction["text"])

    # -----------------------------------------------------
    # 5. Basic counts
    # -----------------------------------------------------

    lines = [
        line
        for line in cleaned.split("\n")
        if line.strip()
    ]
    words = cleaned.split()

    # -----------------------------------------------------
    # 6. Simple extraction fields for frontend
    # -----------------------------------------------------

    detected_dates = detect_dates(cleaned)
    detected_measurements = detect_measurements(cleaned)
    monetary_amounts = detect_money(cleaned)
    key_value_fields = detect_key_values(cleaned)

    relevant_detected_terms = find_relevant_terms(
        cleaned,
        ENERGY_TERMS | WATER_TERMS,
    )

    # -----------------------------------------------------
    # 7. Technical metadata
    # -----------------------------------------------------

    metadata = {
        "file_type": ftype,
        "page_count": extraction["page_count"],
        "word_count": len(words),
        "line_count": len(lines),
        "detected_dates": detected_dates,
        "detected_measurements": detected_measurements,
        "monetary_amounts": monetary_amounts,
        "key_value_fields": key_value_fields,
        "relevant_detected_terms": relevant_detected_terms,
        "extraction_method": extraction["method"],
        "processed_timestamp": datetime.now(
            timezone.utc
        ).isoformat(),
    }

    # -----------------------------------------------------
    # 8. Rich structured JSON
    # -----------------------------------------------------

    structured = build_structured_data(
        file.filename or "document",
        ftype,
        cleaned,
        metadata,
    )

    # -----------------------------------------------------
    # 9. SHA-256 file hash
    # -----------------------------------------------------

    sha256 = hashlib.sha256(buf).hexdigest()

    # -----------------------------------------------------
    # 10. Final response
    # -----------------------------------------------------

    return {
        "document_id": document_id,
        "extracted_text": cleaned,
        "structured_data": structured,
        "metadata": metadata,
        "confidence": (
            0.92
            if "OCR" not in extraction["method"]
            else 0.78
        ),
        "sha256": sha256,
        "extraction_method": extraction["method"],
    }


# =========================================================
# HEALTH CHECK
# =========================================================

@app.get("/health")
async def health():
    """Health check for the processor service."""

    return {
        "status": "ok",
        "service": "greentrust-processor",
    }
